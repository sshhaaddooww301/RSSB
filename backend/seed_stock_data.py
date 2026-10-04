"""
Seed script to import stock items into database (Supabase / PostgreSQL and Local Excel).
"""

import sys
from pathlib import Path
from datetime import datetime, timezone
import sqlalchemy as sa

sys.path.insert(0, str(Path(__file__).parent))

from config import settings
from services.postgres_client import postgres_client, _get_clean_db_url
from services.excel_local import initialize_local_excel_if_needed, read_table, DB_PATH
import openpyxl

RAW_ITEMS = [
    {"Item_No": "6745", "Item_Name": "Atta Wheat", "Unit": "QTL", "Opening_Qty": 280, "Langar_Qty": 240, "Minimum_Stock": 20, "Critical_Stock": 5},
    {"Item_No": "6738", "Item_Name": "Rice Kolam", "Unit": "QTL", "Opening_Qty": 210, "Langar_Qty": 210, "Minimum_Stock": 20, "Critical_Stock": 5},
    {"Item_No": "6744", "Item_Name": "Besan", "Unit": "QTL", "Opening_Qty": 22, "Langar_Qty": 4.4, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "70", "Item_Name": "Dal Rajmah Chitra", "Unit": "QTL", "Opening_Qty": 12, "Langar_Qty": 12, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "48", "Item_Name": "Dal Chana", "Unit": "QTL", "Opening_Qty": 9, "Langar_Qty": 9, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "6685", "Item_Name": "Dal Urad (Mah) Chhilka", "Unit": "QTL", "Opening_Qty": 14, "Langar_Qty": 14, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "11895", "Item_Name": "Dal Urad (Mah) Sabut", "Unit": "QTL", "Opening_Qty": 16, "Langar_Qty": 16, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "65", "Item_Name": "Dal Moong Sabut", "Unit": "QTL", "Opening_Qty": 10, "Langar_Qty": 10, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "49", "Item_Name": "Dal Moong Chhilka", "Unit": "QTL", "Opening_Qty": 12, "Langar_Qty": 12, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "50", "Item_Name": "Dal Moong Dhuli", "Unit": "QTL", "Opening_Qty": 10, "Langar_Qty": 10, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "66", "Item_Name": "Dal Massar Sabut", "Unit": "QTL", "Opening_Qty": 5, "Langar_Qty": 5, "Minimum_Stock": 3, "Critical_Stock": 1},
    {"Item_No": "62", "Item_Name": "Dal Massar Malika", "Unit": "QTL", "Opening_Qty": 16, "Langar_Qty": 16, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "47", "Item_Name": "Dal Arhar", "Unit": "QTL", "Opening_Qty": 3.5, "Langar_Qty": 3.5, "Minimum_Stock": 2, "Critical_Stock": 1},
    {"Item_No": "67", "Item_Name": "Dal Safaid Chana", "Unit": "QTL", "Opening_Qty": 10, "Langar_Qty": 9.84, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "6650", "Item_Name": "Dhania Powder", "Unit": "KG", "Opening_Qty": 200, "Langar_Qty": 160, "Minimum_Stock": 20, "Critical_Stock": 5},
    {"Item_No": "6678", "Item_Name": "Lal Mirch Powder", "Unit": "KG", "Opening_Qty": 100, "Langar_Qty": 82, "Minimum_Stock": 15, "Critical_Stock": 5},
    {"Item_No": "6651", "Item_Name": "Dhania Sabut", "Unit": "KG", "Opening_Qty": 8, "Langar_Qty": 5, "Minimum_Stock": 2, "Critical_Stock": 1},
    {"Item_No": "6673", "Item_Name": "Lal Mirch Sabut", "Unit": "KG", "Opening_Qty": 2, "Langar_Qty": 2, "Minimum_Stock": 1, "Critical_Stock": 0.5},
    {"Item_No": "6668", "Item_Name": "Haldi Powder", "Unit": "KG", "Opening_Qty": 100, "Langar_Qty": 84, "Minimum_Stock": 15, "Critical_Stock": 5},
    {"Item_No": "6771", "Item_Name": "Hing (LS)", "Unit": "KG", "Opening_Qty": 5, "Langar_Qty": 5, "Minimum_Stock": 1, "Critical_Stock": 0.5},
    {"Item_No": "11882", "Item_Name": "Salt Iodised", "Unit": "KG", "Opening_Qty": 1250, "Langar_Qty": 998, "Minimum_Stock": 100, "Critical_Stock": 25},
    {"Item_No": "6652", "Item_Name": "Jeera", "Unit": "KG", "Opening_Qty": 105, "Langar_Qty": 90, "Minimum_Stock": 15, "Critical_Stock": 5},
    {"Item_No": "6672", "Item_Name": "Kali Mirch, Pack of 1 Kg.", "Unit": "KG", "Opening_Qty": 4, "Langar_Qty": 3, "Minimum_Stock": 1, "Critical_Stock": 0.5},
    {"Item_No": "6061", "Item_Name": "Kali Mirch Sabat 100Gms Make: Aplus / The Omegro Farms", "Unit": "PKT", "Opening_Qty": 10, "Langar_Qty": 10, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "6664", "Item_Name": "Kasuri Methi", "Unit": "KG", "Opening_Qty": 15, "Langar_Qty": 7, "Minimum_Stock": 3, "Critical_Stock": 1},
    {"Item_No": "6675", "Item_Name": "Degi Mirch Powder", "Unit": "KG", "Opening_Qty": 31, "Langar_Qty": 20, "Minimum_Stock": 5, "Critical_Stock": 2},
    {"Item_No": "6662", "Item_Name": "Ambchoor", "Unit": "KG", "Opening_Qty": 18, "Langar_Qty": 5, "Minimum_Stock": 3, "Critical_Stock": 1},
    {"Item_No": "6661", "Item_Name": "Ajwain", "Unit": "KG", "Opening_Qty": 40, "Langar_Qty": 4, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "18848", "Item_Name": "Seed Kasuri Methi, Make: Local", "Unit": "KG", "Opening_Qty": 2, "Langar_Qty": 2, "Minimum_Stock": 1, "Critical_Stock": 0.5},
    {"Item_No": "11879", "Item_Name": "Kala Namak", "Unit": "KG", "Opening_Qty": 65, "Langar_Qty": 33, "Minimum_Stock": 10, "Critical_Stock": 2},
    {"Item_No": "6649", "Item_Name": "Illachi Choti", "Unit": "KG", "Opening_Qty": 11, "Langar_Qty": 2, "Minimum_Stock": 2, "Critical_Stock": 1},
    {"Item_No": "6666", "Item_Name": "Garam Masala", "Unit": "KG", "Opening_Qty": 25, "Langar_Qty": 21, "Minimum_Stock": 5, "Critical_Stock": 2},
    {"Item_No": "38469", "Item_Name": "Kalaunji 100Gms, Make: Aplus / The Omegro Farms", "Unit": "PKT", "Opening_Qty": 10, "Langar_Qty": 9, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "6654", "Item_Name": "Saunf", "Unit": "KG", "Opening_Qty": 25, "Langar_Qty": 2, "Minimum_Stock": 3, "Critical_Stock": 1},
    {"Item_No": "6784", "Item_Name": "Rayee", "Unit": "KG", "Opening_Qty": 30, "Langar_Qty": 8, "Minimum_Stock": 5, "Critical_Stock": 1},
    {"Item_No": "15321", "Item_Name": "Poha, Pack of 30Kg, Make: Standard Make", "Unit": "KG", "Opening_Qty": 1050, "Langar_Qty": 50, "Minimum_Stock": 100, "Critical_Stock": 20},
    {"Item_No": "6779", "Item_Name": "Sugar S-30", "Unit": "QTL", "Opening_Qty": 30, "Langar_Qty": 12, "Minimum_Stock": 5, "Critical_Stock": 2},
]


def seed_database():
    now = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    print(f"[*] Seeding {len(RAW_ITEMS)} items...")

    # 1. Seed Supabase / PostgreSQL
    if postgres_client.is_configured():
        print("[*] Connecting to Supabase / PostgreSQL...")
        try:
            postgres_client.initialize_schema_if_needed()
            engine = postgres_client._get_engine()
            with engine.connect() as conn:
                for item in RAW_ITEMS:
                    item_no = item["Item_No"]
                    stmt = sa.text("""
                        INSERT INTO items (
                            "Item_ID", "Item_No", "Item_Name", "SKU", "Unit",
                            "Opening_Qty", "Langar_Qty", "Minimum_Stock", "Critical_Stock",
                            "Current_Stock", "Status", "Created_Date", "Updated_Date"
                        ) VALUES (
                            :item_id, :item_no, :item_name, :sku, :unit,
                            :opening_qty, :langar_qty, :min_stock, :crit_stock,
                            :current_stock, 'Active', :created_date, :updated_date
                        )
                        ON CONFLICT ("Item_No") DO UPDATE SET
                            "Item_Name" = EXCLUDED."Item_Name",
                            "SKU" = EXCLUDED."SKU",
                            "Unit" = EXCLUDED."Unit",
                            "Opening_Qty" = EXCLUDED."Opening_Qty",
                            "Langar_Qty" = EXCLUDED."Langar_Qty",
                            "Minimum_Stock" = EXCLUDED."Minimum_Stock",
                            "Critical_Stock" = EXCLUDED."Critical_Stock",
                            "Current_Stock" = EXCLUDED."Current_Stock",
                            "Updated_Date" = EXCLUDED."Updated_Date";
                    """)
                    conn.execute(stmt, {
                        "item_id": item_no,
                        "item_no": item_no,
                        "item_name": item["Item_Name"],
                        "sku": f"{item['Langar_Qty']} {item['Unit']}",
                        "unit": item["Unit"],
                        "opening_qty": item["Opening_Qty"],
                        "langar_qty": item["Langar_Qty"],
                        "min_stock": item["Minimum_Stock"],
                        "crit_stock": item["Critical_Stock"],
                        "current_stock": item["Opening_Qty"],
                        "created_date": now,
                        "updated_date": now,
                    })
                conn.commit()
            print("[+] Successfully seeded/upserted Supabase / PostgreSQL items table!")
        except Exception as e:
            print(f"[-] Supabase seed error: {e}")

    # 2. Seed Local Excel Database
    initialize_local_excel_if_needed()
    wb = openpyxl.load_workbook(DB_PATH)
    if "ITEMS" in wb.sheetnames:
        ws = wb["ITEMS"]
        # Clear existing rows keeping header
        while ws.max_row > 1:
            ws.delete_rows(2, 1)

        for item in RAW_ITEMS:
            item_no = item["Item_No"]
            row_values = [
                item_no,
                item_no,
                item["Item_Name"],
                f"{item['Langar_Qty']} {item['Unit']}",
                item["Unit"],
                item["Opening_Qty"],
                item["Langar_Qty"],
                item["Minimum_Stock"],
                item["Critical_Stock"],
                item["Opening_Qty"],
                "Active",
                now,
                now,
            ]
            ws.append(row_values)
        wb.save(DB_PATH)
        print("[+] Successfully updated local Excel Kitchen_Stock_Database.xlsx with all 37 items!")


if __name__ == "__main__":
    seed_database()
